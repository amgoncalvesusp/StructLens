"""Final-round RED contracts for Task 13 provenance and scientific views.

These tests intentionally describe the last rejected behaviors.  They use real
report-service output and immutable snapshots; doubles are narrow command,
executor, dialog, and writer spies at the side-effect boundaries only.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication

from structlens.application.chart_data import ChartDataset, ChartSeries
from structlens.application.chart_export import export_chart_xlsx
from structlens.core.errors import AnalysisCancelledError, ProjectSchemaError
from structlens.core.reports import AnalysisReport
from structlens.integrations.pymol.adapter import PyMOLAdapter
from structlens.plugin.gui import qt_panel
from structlens.plugin.gui.main_panel import build_qt_panel
from structlens.plugin.gui.model import CanonicalReportBinding
from structlens.plugin.gui.project_transaction import save_canonical_project, stage_project
from structlens.plugin.visualization.renderer import VisualizationState

from ._task13_round4_fixtures import rich_report
from ._task13_round5_fixtures import (
    gapped_deviation_report,
    multi_selected_report,
    typed_report,
    unavailable_typed_report,
)


@pytest.fixture(scope="module")
def application() -> QApplication:
    app = QApplication.instance() or QApplication([])
    yield app
    app.quit()


@pytest.fixture()
def controller(application):
    panel = build_qt_panel(command=None)
    value = panel._structlens_controller
    yield value
    value.close()
    panel.deleteLater()
    application.processEvents()


def _install_report(controller: object, request: object, report: AnalysisReport) -> None:
    """Install a real report through the panel's existing source boundary."""

    assert hasattr(request, "reference_selection")
    assert hasattr(request, "target_selection")
    reference_path = Path(request.reference_selection.path)  # type: ignore[attr-defined]
    target_path = Path(request.target_selection.path)  # type: ignore[attr-defined]
    assert controller._load_source("reference", reference_path)  # type: ignore[attr-defined]
    assert controller._load_source("target", target_path)  # type: ignore[attr-defined]
    controller._report_request = request  # type: ignore[attr-defined]
    controller._analysis_finished(report)  # type: ignore[attr-defined]


class _ObjectCommand:
    """Small PyMOL object proxy that records bytes and object identity."""

    def __init__(self, content: bytes, names: tuple[str, ...] = ("complex",)) -> None:
        self.content = content
        self.names = list(names)
        self.saved: list[tuple[str, str]] = []
        self.loaded: list[tuple[str, str, bytes]] = []
        self.calls: list[tuple[str, tuple[object, ...]]] = []

    def get_names(self, kind: str) -> list[str]:
        assert kind == "objects"
        return list(self.names)

    def save(self, path: str, name: str) -> None:
        self.saved.append((path, name))
        Path(path).write_bytes(self.content)

    def load(self, path: str, name: str) -> None:
        payload = Path(path).read_bytes()
        self.loaded.append((path, name, payload))
        if name not in self.names:
            self.names.append(name)

    def delete(self, name: str) -> None:
        self.calls.append(("delete", (name,)))
        self.names = [value for value in self.names if value != name]

    def __getattr__(self, name: str):
        def record(*args: object) -> None:
            self.calls.append((name, args))

        return record


def test_selected_pymol_object_is_re_materialized_as_unique_snapshot_owned_object(controller, monkeypatch) -> None:
    """A user object is capture input only; canonical state must own a new object."""

    fixture = Path("tests/fixtures/parsing/numbering_altloc.pdb").read_bytes()
    command = _ObjectCommand(fixture, names=("complex",))
    controller.command = command
    monkeypatch.setattr(controller.w.QInputDialog, "getItem", lambda *_: ("complex", True))

    controller._use_pymol_object("reference")

    assert controller.reference_loaded_source is not None
    assert controller.reference_object_name not in {None, "complex"}
    assert controller.reference_object_name.startswith("structlens_")
    assert command.loaded, "captured bytes must be loaded into a StructLens-owned object"
    assert command.loaded[-1][1] == controller.reference_object_name
    assert command.loaded[-1][2] == fixture

    # Deleting or mutating the original PyMOL object after capture cannot
    # invalidate the immutable source that will feed canonical operations.
    command.names = [name for name in command.names if name != "complex"]
    command.content = b"ATOM      1  CA  BAD X   1       0.000   0.000   0.000\nEND\n"
    assert controller.reference_loaded_source.snapshot.decompressed_bytes == fixture
    assert controller.reference_structure is not None
    assert controller.reference_structure.chains


def test_canonical_pymol_report_uses_snapshot_owned_model_two_chain_b_objects(tmp_path: Path) -> None:
    """Apply must constrain model 2 / chain B and never reuse mutable names."""

    request, report = multi_selected_report(tmp_path)
    assert report.analysis is not None
    command = _ObjectCommand(request.reference_snapshot.decompressed_bytes, names=("same-basename",))
    adapter = PyMOLAdapter(command, project_id="round5")
    apply_report = getattr(adapter, "apply_report", None)
    assert callable(apply_report), "canonical PyMOL adapter needs a snapshot-bound report operation"

    # The source paths are deliberately destroyed before visualization.  The
    # report remains valid because snapshots, not paths, are its authority.
    Path(request.reference_selection.path).unlink()
    Path(request.target_selection.path).unlink()
    apply_report(
        report,
        snapshots=(request.reference_snapshot, request.target_snapshot),
        state=VisualizationState(),
        reference_object="same-basename",
        target_object="same-basename",
    )

    loaded_names = [name for _path, name, _bytes in command.loaded]
    assert len(loaded_names) >= 2
    assert len(set(loaded_names)) == len(loaded_names)
    assert all(name != "same-basename" and name.startswith("structlens_") for name in loaded_names)
    assert command.loaded[0][2] == request.reference_snapshot.decompressed_bytes
    assert command.loaded[1][2] == request.target_snapshot.decompressed_bytes
    create_calls = [args for method, args in command.calls if method == "create" and len(args) == 4]
    assert len(create_calls) >= 2
    assert create_calls[0][2] == 2
    assert create_calls[0][3] == 1
    assert create_calls[1][2] == 1
    assert create_calls[1][3] == 1
    assert all("chain B" in str(args[1]) for args in create_calls[:2])
    select_calls = [args for method, args in command.calls if method == "select"]
    assert select_calls
    assert all("same-basename" not in str(args[1]) for args in select_calls)
    assert all("structlens_" in str(args[1]) for args in select_calls)


class _DeferredFuture:
    def __init__(self, result: object) -> None:
        self._result = result

    def done(self) -> bool:
        return True

    def result(self) -> object:
        if isinstance(self._result, BaseException):
            raise self._result
        return self._result

    def cancel(self) -> bool:
        return False


class _DeferredExecutor:
    def __init__(self, result: object) -> None:
        self.result = result

    def submit(self, _function: object, _request: object) -> _DeferredFuture:
        return _DeferredFuture(self.result)


def test_failed_pending_comparison_preserves_accepted_binding_for_export_and_pymol(
    controller, monkeypatch, tmp_path: Path
) -> None:
    """A failed next request must not replace the accepted canonical request."""

    (tmp_path / "accepted").mkdir()
    accepted_request, accepted_report = rich_report(tmp_path / "accepted")
    _install_report(controller, accepted_request, accepted_report)
    accepted_binding = controller.model.binding
    assert accepted_binding is not None

    (tmp_path / "pending").mkdir()
    pending_request, _ = rich_report(tmp_path / "pending")
    monkeypatch.setattr(
        "structlens.plugin.gui.qt_sources_mixin._ANALYSIS_EXECUTOR",
        _DeferredExecutor(AnalysisCancelledError()),
    )
    controller._submit_report_request(pending_request)
    controller._poll_analysis()

    assert controller.model.binding is accepted_binding
    assert controller.model.report is accepted_report
    assert controller._report_request is accepted_request
    assert controller._canonical_binding() is accepted_binding
    exported: list[tuple[object, object, object]] = []
    monkeypatch.setattr(
        controller.w.QFileDialog,
        "getSaveFileName",
        lambda *_args: (str(tmp_path / "accepted.json"), "JSON"),
    )
    controller._export(
        "json",
        lambda value, path, **kwargs: exported.append((value, path, kwargs)),
        "JSON",
    )
    assert exported and exported[0][0] is accepted_report
    assert exported[0][2]["snapshots"] == accepted_binding.snapshots


def test_mismatched_service_report_is_rejected_before_any_model_or_presentation_replacement(
    controller, tmp_path: Path
) -> None:
    """A result for another request must leave the complete accepted state intact."""

    (tmp_path / "accepted").mkdir()
    accepted_request, accepted_report = rich_report(tmp_path / "accepted")
    _install_report(controller, accepted_request, accepted_report)
    before_model = controller.model
    before_presentation = controller._report_presentation
    before_history = controller._report_history
    before_request = controller._report_request
    errors: list[str] = []
    controller._show_error = errors.append

    (tmp_path / "stale").mkdir()
    stale_request, stale_report = multi_selected_report(tmp_path / "stale")
    assert stale_report.report_id != accepted_report.report_id
    controller._report_request = accepted_request
    controller._analysis_finished(stale_report)

    assert errors
    assert controller.model is before_model
    assert controller.model.report is accepted_report
    assert controller.model.binding is before_model.binding
    assert controller._report_presentation is before_presentation
    assert controller._report_history == before_history
    assert controller._report_request is before_request


def test_canonical_open_commits_staged_binding_directly_without_reexecuting_report(
    controller, monkeypatch, tmp_path: Path
) -> None:
    """Open must consume the validated candidate, not the controller's old request."""

    (tmp_path / "saved").mkdir()
    request, report = rich_report(tmp_path / "saved")
    binding = CanonicalReportBinding.create(report, request)
    project_path = tmp_path / "saved.structlens.json"
    save_canonical_project(binding, project_path, visualization_state={})
    candidate = stage_project(project_path)
    assert hasattr(candidate, "binding")

    (tmp_path / "current").mkdir()
    current_request, current_report = typed_report(tmp_path / "current")
    _install_report(controller, current_request, current_report)
    calls: list[object] = []
    monkeypatch.setattr(controller, "_analysis_finished", lambda value: calls.append(value))

    controller._commit_canonical_project(candidate, project_path)

    assert calls == []
    assert controller.model.binding is candidate.binding
    assert controller.model.report is candidate.binding.report
    assert controller._report_request is candidate.binding.request


def test_malformed_visualization_state_rolls_back_populated_panel_before_commit(
    controller, monkeypatch, tmp_path: Path
) -> None:
    """Unsafe persisted controls must fail during staging and leave all UI state untouched."""

    (tmp_path / "saved").mkdir()
    request, report = rich_report(tmp_path / "saved")
    project_path = tmp_path / "malformed.structlens.json"
    save_canonical_project(CanonicalReportBinding.create(report, request), project_path, visualization_state={})
    payload = json.loads(project_path.read_text(encoding="utf-8"))
    payload["visualization_state"] = {"highlight_filter": "not-a-filter"}
    project_path.write_text(json.dumps(payload), encoding="utf-8")

    (tmp_path / "current").mkdir()
    current_request, current_report = rich_report(tmp_path / "current")
    _install_report(controller, current_request, current_report)
    before_model = controller.model
    before_source = controller.reference_loaded_source
    before_visualization = controller.model.visualization_state
    errors: list[str] = []
    controller._show_error = errors.append
    monkeypatch.setattr(
        controller.w.QFileDialog,
        "getOpenFileName",
        lambda *_args: (str(project_path), "StructLens project (*.json)"),
    )

    controller._open_project()

    assert errors
    assert controller.model is before_model
    assert controller.reference_loaded_source is before_source
    assert controller.model.visualization_state == before_visualization


@pytest.mark.parametrize("action", ("xlsx", "image"))
def test_canonical_chart_export_verifies_binding_before_dialog_or_writer(
    controller, monkeypatch, tmp_path: Path, action: str
) -> None:
    """Every chart path must reject stale provenance before any side effect."""

    (tmp_path / "report").mkdir()
    request, report = rich_report(tmp_path / "report")
    _install_report(controller, request, report)
    controller.set_chart_datasets(
        {
            "MSA conservation profile": ChartDataset(
                "msa",
                "MSA",
                "Alignment column",
                "Conservation",
                "fraction",
                (ChartSeries("sequence", ((1.0, 1.0),)),),
                "authoritative",
            )
        }
    )
    controller.chart_combo.setCurrentText("MSA conservation profile")
    (tmp_path / "stale").mkdir()
    stale_request, _stale_report = rich_report(tmp_path / "stale")
    controller._report_request = stale_request
    dialogs: list[object] = []
    writers: list[object] = []
    errors: list[str] = []
    controller._show_error = errors.append
    monkeypatch.setattr(
        controller.w.QFileDialog,
        "getSaveFileName",
        lambda *_args: dialogs.append(True) or (str(tmp_path / "out"), "output"),
    )
    if action == "xlsx":
        monkeypatch.setattr(qt_panel, "export_chart_xlsx", lambda *_args, **_kwargs: writers.append(True))
        controller._export_chart_xlsx()
    else:
        monkeypatch.setattr(qt_panel, "export_chart_image", lambda *_args, **_kwargs: writers.append(True))
        controller._export_chart_image("png", 300)

    assert not dialogs
    assert not writers
    assert errors and ("stale" in errors[-1].lower() or "verified" in errors[-1].lower())


@pytest.mark.parametrize("action", ("xlsx", "image"))
def test_render_only_chart_export_fails_closed_before_dialog(
    controller, monkeypatch, tmp_path: Path, action: str
) -> None:
    """A presentation-only compatibility report is not canonical chart evidence."""

    _request, report = rich_report(tmp_path)
    controller.model = controller.model.with_report(report)
    controller.set_chart_datasets(
        {
            "MSA conservation profile": ChartDataset(
                "msa", "MSA", "column", "conservation", "fraction", (), "authoritative"
            )
        }
    )
    controller.chart_combo.setCurrentText("MSA conservation profile")
    dialogs: list[object] = []
    errors: list[str] = []
    controller._show_error = errors.append
    monkeypatch.setattr(
        controller.w.QFileDialog,
        "getSaveFileName",
        lambda *_args: dialogs.append(True) or (str(tmp_path / "out"), "output"),
    )

    if action == "xlsx":
        controller._export_chart_xlsx()
    else:
        controller._export_chart_image("png", 300)

    assert not dialogs
    assert errors and "verified" in errors[-1].lower()


def test_malformed_visualization_state_is_rejected_during_canonical_staging(tmp_path: Path) -> None:
    """Persisted visualization controls must be validated before GUI mutation."""

    (tmp_path / "saved").mkdir()
    request, report = rich_report(tmp_path / "saved")
    project_path = tmp_path / "malformed-staging.structlens.json"
    save_canonical_project(CanonicalReportBinding.create(report, request), project_path, visualization_state={})
    payload = json.loads(project_path.read_text(encoding="utf-8"))
    payload["visualization_state"] = {"highlight_filter": "not-a-filter"}
    project_path.write_text(json.dumps(payload), encoding="utf-8")

    with pytest.raises((ValueError, TypeError, ProjectSchemaError), match="visualization|highlight|filter"):
        stage_project(project_path)


def test_structural_deviation_chart_preserves_authoritative_residue_labels_and_alignment_metadata(
    controller, tmp_path: Path
) -> None:
    """Gaps, insertion codes, and numbering are never replaced by 1,2,3."""

    _request, report = gapped_deviation_report(tmp_path)
    controller._analysis_finished(report)
    dataset = controller._chart_datasets["Structural deviation profile"]
    assert isinstance(dataset, ChartDataset)
    series = dataset.series[0]
    assert series.labels[:3] == ("A:10", "A:21A", "A:42")
    assert series.points[:3] == ((10.0, series.points[0][1]), (21.0, series.points[1][1]), (42.0, series.points[2][1]))
    assert tuple(item["alignment_index"] for item in series.metadata[:3]) == (2, 7, 11)
    assert series.points[:3] != tuple((float(index + 1), value) for index, (_x, value) in enumerate(series.points[:3]))


def test_chart_axes_status_and_xlsx_metadata_include_dataset_units(controller, monkeypatch, tmp_path: Path) -> None:
    """A numeric chart without an explicit unit is scientifically ambiguous."""

    request, report = rich_report(tmp_path)
    _install_report(controller, request, report)
    controller.chart_combo.setCurrentText("Structural deviation profile")
    dataset = controller._chart_datasets["Structural deviation profile"]
    assert isinstance(dataset, ChartDataset)
    assert dataset.unit == "Å"
    assert "Å" in dataset.y_label
    assert "Å" in controller.chart_preview_status.text()

    captured: list[tuple[object, object]] = []
    monkeypatch.setattr(
        qt_panel,
        "export_chart_xlsx",
        lambda value, path: captured.append((value, path)),
    )
    monkeypatch.setattr(
        controller.w.QFileDialog,
        "getSaveFileName",
        lambda *_args: (str(tmp_path / "chart.xlsx"), "XLSX"),
    )
    controller._export_chart_xlsx()
    assert captured
    exported = captured[0][0]
    assert exported.unit == "Å"
    assert "Å" in exported.y_label


def test_chart_xlsx_serializes_the_dataset_unit_even_when_axis_label_is_plain(tmp_path: Path) -> None:
    """XLSX consumers need units even when a producer uses a plain axis label."""

    dataset = ChartDataset(
        "unit-check",
        "Unit check",
        "Reference residue",
        "Displacement",
        "Å",
        (ChartSeries("target", ((10.0, 1.5),), ("A:10",)),),
        "descriptive",
    )
    output = tmp_path / "unit-check.xlsx"
    export_chart_xlsx(dataset, output)
    sheet = load_workbook(output, read_only=True).active
    assert sheet is not None
    header = tuple(str(cell.value) for cell in next(sheet.iter_rows(max_row=1)))
    assert any("Å" in value for value in header)


def test_visible_typed_rows_preserve_interaction_and_site_values_units_and_reasons(controller, tmp_path: Path) -> None:
    """The Qt view must expose all typed geometry and native unavailable reasons."""

    (tmp_path / "typed").mkdir()
    request, report = typed_report(tmp_path / "typed")
    _install_report(controller, request, report)
    interactions_text = controller.report_section_labels["interactions"].text()
    sites_text = controller.report_section_labels["sites"].text()
    visible_text = " ".join((interactions_text, sites_text, controller.site_status_label.text()))
    for value in ("2.71", "121.5", "1.2", "0.8", "0.5", "2.3", "12.5", "42.0", "0.75", "0.25"):
        assert value in visible_text
    assert "Å" in visible_text
    assert "degrees" in visible_text or "°" in visible_text
    assert "fraction" in visible_text

    (tmp_path / "unavailable").mkdir()
    _request, unavailable = unavailable_typed_report(tmp_path / "unavailable")
    controller._analysis_finished(unavailable)
    assert "report.interactions.failed" in controller.report_section_labels["interactions"].text()
    assert "could not be completed" in controller.report_section_labels["interactions"].text()
    assert "report.sites.invalid" in controller.report_section_labels["sites"].text()
    assert "selected site input is invalid" in controller.report_section_labels["sites"].text()


def test_visualization_service_is_the_only_report_selection_route(controller, monkeypatch, tmp_path: Path) -> None:
    """Normal GUI PyMOL routing must use the typed report selection service."""

    request, report = rich_report(tmp_path)
    _install_report(controller, request, report)
    controller.command = object()
    controller.reference_object_name = "owned-reference"
    controller.target_object_name = "owned-target"
    calls: list[object] = []
    monkeypatch.setattr(controller._visualization_service, "select", lambda *args: calls.append(args) or ())

    class SpyAdapter:
        def __init__(self) -> None:
            self.calls: list[tuple[tuple[object, ...], dict[str, object]]] = []

        def apply_report(self, *args: object, **kwargs: object) -> None:
            self.calls.append((args, kwargs))

        def apply(self, *_args: object, **_kwargs: object) -> None:
            raise AssertionError("canonical visualization must not use the legacy adapter method")

        def reset(self) -> None:
            return None

    adapter = SpyAdapter()
    controller._pymol = adapter

    controller._apply_visualization()

    assert calls
    assert adapter.calls
    assert adapter.calls[0][0][0] is report
    assert adapter.calls[0][1]["snapshots"] == controller.model.binding.snapshots
    assert adapter.calls[0][1]["state"] == controller.model.visualization_state


__all__ = []
