from __future__ import annotations

import hashlib
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook

import structlens.application.report_exports as report_exports
from structlens.application.export_service import (
    export_report_csv,
    export_report_json,
    export_report_tsv,
    export_report_xlsx,
)
from structlens.core.evidence import Availability, Diagnostic, DiagnosticSeverity
from structlens.core.interactions import InteractionRecord, InteractionType
from structlens.core.models import ResidueId
from structlens.core.parsing import InputSelection, SourceSnapshot, StructureFormat
from structlens.core.provenance import MethodProvenance
from structlens.core.quality import StructureQualityReport
from structlens.core.reports import AnalysisReport, InputQualityBundle, SectionAvailability
from structlens.core.sites import SiteMetrics


def _report() -> tuple[AnalysisReport, SourceSnapshot, SourceSnapshot]:
    reference_bytes = b"reference"
    target_bytes = b"target"
    reference_id = hashlib.sha256(reference_bytes).hexdigest()
    target_id = hashlib.sha256(target_bytes).hexdigest()
    reference_snapshot = SourceSnapshot(
        reference_bytes, reference_id, reference_id, reference_id, "reference.pdb", "pdb", False
    )
    target_snapshot = SourceSnapshot(target_bytes, target_id, target_id, target_id, "target.pdb", "pdb", False)
    reference = InputSelection(reference_id, "reference.pdb", StructureFormat.PDB, "1")
    target = InputSelection(target_id, "target.pdb", StructureFormat.PDB, "1")
    report = AnalysisReport(
        reference,
        target,
        InputQualityBundle(
            StructureQualityReport(Availability.NOT_APPLICABLE), StructureQualityReport(Availability.NOT_APPLICABLE)
        ),
        availability=SectionAvailability(),
        provenance=MethodProvenance(
            "test.report",
            "1",
            input_hashes={
                "reference_content": reference.content_id,
                "target_content": target.content_id,
                "reference_raw": reference_snapshot.raw_sha256,
                "target_raw": target_snapshot.raw_sha256,
            },
        ),
    )
    return report, reference_snapshot, target_snapshot


def test_v04_workbook_has_typed_evidence_sheets_and_provenance(tmp_path: Path) -> None:
    path = tmp_path / "report.xlsx"
    report, reference, target = _report()
    export_report_xlsx(report, path, snapshots=(reference, target))
    workbook = load_workbook(path, read_only=True, data_only=True)
    assert {"Summary", "QC", "Pockets", "Matches", "Lining Residues", "Diagnostics", "Methods", "Provenance"}.issubset(
        workbook.sheetnames
    )
    assert workbook["Summary"]["A1"].value == "Metric"
    assert workbook["Provenance"]["A1"].value == "Key"


def test_tidy_exports_keep_status_reason_and_blank_missing_values(tmp_path: Path) -> None:
    csv_path = tmp_path / "pockets.csv"
    tsv_path = tmp_path / "pockets.tsv"
    report, reference, target = _report()
    export_report_csv(report, csv_path, snapshots=(reference, target))
    export_report_tsv(report, tsv_path, snapshots=(reference, target))
    assert "status" in csv_path.read_text(encoding="utf-8").splitlines()[0]
    assert "reason" in tsv_path.read_text(encoding="utf-8").splitlines()[0]


@pytest.mark.parametrize("exporter", (export_report_csv, export_report_xlsx))
def test_exports_reject_excel_cell_strings_before_output_allocation(tmp_path: Path, exporter: object) -> None:
    report, reference, target = _report()
    long_selection = InputSelection(reference.content_id, "x" * 32_768, StructureFormat.PDB, "1")
    oversized = replace(report, reference_selection=long_selection)
    with pytest.raises(ValueError, match="cell|length"):
        exporter(oversized, tmp_path / "oversized.out", snapshots=(reference, target))  # type: ignore[operator]


def test_delimited_exports_reject_row_caps_before_stringio_growth(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    report, reference, target = _report()
    monkeypatch.setattr(report_exports, "_MAX_EXPORT_ROWS", 1)
    with pytest.raises(ValueError, match="row"):
        export_report_csv(report, tmp_path / "capped.csv", snapshots=(reference, target))


def test_rows_preserve_site_interaction_and_report_diagnostic_sections() -> None:
    report, reference, target = _report()
    residue = ResidueId(target.content_id, "1", "A", "10", None, "ALA")
    rich = replace(
        report,
        availability=replace(report.availability, sites=Availability.AVAILABLE, interactions=Availability.AVAILABLE),
        sites=(SiteMetrics("active", target.content_id, 1, 1.0),),
        interactions=(
            InteractionRecord(target.content_id, InteractionType.HYDROPHOBIC, residue, None, "CA", None, 4.0),
        ),
        diagnostics=(Diagnostic("EXPORT_NOTE", DiagnosticSeverity.INFO, "exported"),),
    )
    rows = report_exports._rows(rich)
    assert {row["section"] for row in rows} >= {"site_metrics", "matches", "diagnostics"}


def test_json_export_writes_verified_canonical_payload(tmp_path: Path) -> None:
    report, reference, target = _report()
    path = tmp_path / "report.json"
    export_report_json(report, path, snapshots=(reference, target))
    assert '"report_id"' in path.read_text(encoding="utf-8")


def test_export_helpers_fail_closed_at_row_cell_and_size_bounds(monkeypatch: pytest.MonkeyPatch) -> None:
    assert report_exports._safe("=formula") == "'=formula"
    assert report_exports._safe(3) == 3
    assert report_exports._provenance_hash({"reference-source-raw": "digest"}, "reference", "raw") == "digest"
    assert report_exports._provenance_hash({}, "reference", "raw") is None
    monkeypatch.setattr(report_exports, "_MAX_CELL_STRING", 2)
    with pytest.raises(ValueError, match="cell"):
        report_exports._json("long")
    row = {key: "value" for key in report_exports._FIELDS}
    monkeypatch.setattr(report_exports, "_MAX_EXPORT_ROWS", 0)
    with pytest.raises(ValueError, match="row"):
        report_exports._validate_rows([row])
    monkeypatch.setattr(report_exports, "_MAX_EXPORT_ROWS", 100_000)
    monkeypatch.setattr(report_exports, "_MAX_EXPORT_CELLS", 1)
    with pytest.raises(ValueError, match="cell"):
        report_exports._validate_rows([row])
    monkeypatch.setattr(report_exports, "_MAX_CELL_STRING", 32_767)
    monkeypatch.setattr(report_exports, "_MAX_EXPORT_CELLS", 1_000_000)
    monkeypatch.setattr(report_exports, "_MAX_EXPORT_BYTES", 1)
    with pytest.raises(ValueError, match="serialized"):
        report_exports._validate_rows([row])
