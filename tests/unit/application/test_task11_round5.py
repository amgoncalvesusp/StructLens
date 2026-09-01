from __future__ import annotations

import csv
import gzip
import hashlib
import importlib
import json
import os
from dataclasses import replace
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

import structlens.application.report_exports as report_exports
import structlens.core.reports.safe_io as safe_io
import structlens.core.reports.schema as schema_module
from structlens.application.export_service import export_report_csv, export_report_tsv, export_report_xlsx
from structlens.application.report_serialization import report_to_dict, serialize_report
from structlens.application.report_snapshot_io import load_snapshot, snapshot_path, store_snapshot
from structlens.core.evidence import (
    Availability,
    Diagnostic,
    DiagnosticSeverity,
    PocketEvidenceChannel,
    PocketEvidenceSections,
)
from structlens.core.models import ResidueId
from structlens.core.parsing import SourceSnapshot, capture_snapshot
from structlens.core.pockets.matching import PocketMatchingSettings
from structlens.core.provenance import MethodProvenance
from structlens.core.reports.schema import validate_analysis_report_payload

fixtures = importlib.import_module("tests.unit.application.test_task11_round1")
round2 = importlib.import_module("tests.unit.application.test_task11_round2")


def _records(path: Path, delimiter: str = ",") -> list[dict[str, str]]:
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle, delimiter=delimiter))


def _sheet_rows(path: Path, name: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    values = list(workbook[name].iter_rows(values_only=True))
    headers = tuple(str(value) for value in values[0])
    return [dict(zip(headers, row, strict=True)) for row in values[1:]]


def _report_with_insertion_code() -> tuple[Any, SourceSnapshot, SourceSnapshot]:
    report, reference, target, _event = round2._mutation_report()
    assert report.pockets is not None
    residue = ResidueId("reference", "2", "A", "10", "B", "GLY")
    lining = replace(report.pockets.lining_residues[0], residues=(residue,))
    return replace(report, pockets=replace(report.pockets, lining_residues=(lining,))), reference, target


def test_pocket_tables_retain_complete_candidate_matching_and_residue_evidence(tmp_path: Path) -> None:
    report, reference, target = _report_with_insertion_code()
    assert report.pockets is not None and report.pockets.matching is not None
    candidate = report.pockets.detections[0].candidates[0]
    expected_residue = report.pockets.lining_residues[0].to_json()["residues"][0]

    csv_path = tmp_path / "complete.csv"
    xlsx_path = tmp_path / "complete.xlsx"
    export_report_csv(report, csv_path, snapshots=(reference, target))
    export_report_xlsx(report, xlsx_path, snapshots=(reference, target))

    records = _records(csv_path)
    candidate_row = next(row for row in records if row["metric"] == "candidate_evidence")
    matching_row = next(row for row in records if row["metric"] == "matching_evidence")
    residue_row = next(row for row in records if row["section"] == "lining_residues" and row["metric"] == "residue")
    assert json.loads(candidate_row["value"]) == candidate.to_json()
    assert json.loads(matching_row["value"]) == report.pockets.matching.to_json()
    assert json.loads(residue_row["value"]) == expected_residue
    assert json.loads(candidate_row["value"])["alpha_spheres"][0]["touching_atom_ids"] == ["a", "b", "c", "d"]
    assert json.loads(candidate_row["value"])["lineage"] == {
        "selection_id": candidate.selection_id,
        "source_content_id": candidate.source_content_id,
    }

    pocket_rows = _sheet_rows(xlsx_path, "Pockets")
    match_rows = _sheet_rows(xlsx_path, "Matches")
    lining_rows = _sheet_rows(xlsx_path, "Lining Residues")
    assert (
        json.loads(next(row["Value"] for row in pocket_rows if row["Metric"] == "candidate_evidence"))
        == candidate.to_json()
    )
    assert (
        json.loads(next(row["value"] for row in match_rows if row["metric"] == "matching_evidence"))
        == report.pockets.matching.to_json()
    )
    assert json.loads(next(row["value"] for row in lining_rows if row["metric"] == "residue")) == expected_residue


def test_empty_matching_still_exports_settings_transform_scores_and_diagnostics(tmp_path: Path) -> None:
    report, reference, target, _event = round2._mutation_report()
    assert report.pockets is not None and report.pockets.matching is not None
    diagnostic = Diagnostic("matching.empty", DiagnosticSeverity.INFO, "No candidate pairs")
    matching = replace(
        report.pockets.matching,
        availability=Availability.NOT_APPLICABLE,
        matches=(),
        diagnostics=(diagnostic,),
        settings=PocketMatchingSettings(minimum_lining_jaccard=0.25),
        optimal_score=0.0,
        second_best_score=0.125,
    )
    report = replace(report, pockets=replace(report.pockets, matching=matching))
    path = tmp_path / "empty-matching.xlsx"
    export_report_xlsx(report, path, snapshots=(reference, target))

    rows = {row["metric"]: row for row in report_exports._rows(report) if row["section"] == "pocket_matches"}
    assert json.loads(rows["matching_settings"]["value"]) == matching.settings.to_json()
    assert json.loads(rows["matching_transform"]["value"]) == matching.to_json()["transform"]
    assert json.loads(rows["matching_diagnostics"]["value"]) == [diagnostic.to_json()]
    assert rows["matching_optimal_score"]["value"] == 0.0
    assert rows["matching_second_best_score"]["value"] == 0.125
    workbook_rows = _sheet_rows(path, "Matches")
    workbook_metrics = {row["metric"]: row["value"] for row in workbook_rows}
    assert json.loads(workbook_metrics["matching_settings"]) == matching.settings.to_json()
    assert workbook_metrics["matching_optimal_score"] == 0.0
    assert workbook_metrics["matching_second_best_score"] == 0.125


def test_tabular_exports_retain_report_units_and_every_native_concordance_diagnostic(tmp_path: Path) -> None:
    report, reference, target, _event = round2._mutation_report()
    assert report.pockets is not None
    statuses = (
        Availability.INVALID_INPUT,
        Availability.DEPENDENCY_UNAVAILABLE,
        Availability.NUMERICAL_FAILURE,
        Availability.NOT_DETECTED,
        Availability.NOT_APPLICABLE,
        Availability.AVAILABLE,
        Availability.INVALID_INPUT,
    )
    channels: dict[str, PocketEvidenceChannel] = {}
    for index, (name, status) in enumerate(
        zip(
            (
                "geometry",
                "ligand_support",
                "parameter_persistence",
                "volume_sensitivity",
                "match_ambiguity",
                "qc",
                "interactions",
            ),
            statuses,
            strict=True,
        )
    ):
        channels[name] = PocketEvidenceChannel(
            status,
            measure={"native": index},
            units={"zeta": "angstrom", "alpha": "fraction"} if index == 0 else {"native": "count"},
            diagnostics=(
                Diagnostic(
                    f"concordance.{name}",
                    DiagnosticSeverity.WARNING,
                    f"Native {name} diagnostic",
                    source_id=f"source-{name}",
                ),
            ),
            provenance=MethodProvenance(
                f"method.{name}",
                "1",
                parameters={"ordinal": index},
                input_hashes={"source": reference.content_id},
            ).to_json(),
        )
    pockets = replace(
        report.pockets,
        concordance=PocketEvidenceSections(**channels),
        units={"zeta_volume": "angstrom^3", "alpha_score": "fraction"},
    )
    rich = replace(report, pockets=pockets)

    csv_paths = (tmp_path / "rich-a.csv", tmp_path / "rich-b.csv")
    tsv_paths = (tmp_path / "rich-a.tsv", tmp_path / "rich-b.tsv")
    xlsx_paths = (tmp_path / "rich-a.xlsx", tmp_path / "rich-b.xlsx")
    for path in csv_paths:
        export_report_csv(rich, path, snapshots=(reference, target))
    for path in tsv_paths:
        export_report_tsv(rich, path, snapshots=(reference, target))
    for path in xlsx_paths:
        export_report_xlsx(rich, path, snapshots=(reference, target))
    assert csv_paths[0].read_bytes() == csv_paths[1].read_bytes()
    assert tsv_paths[0].read_bytes() == tsv_paths[1].read_bytes()
    assert xlsx_paths[0].read_bytes() == xlsx_paths[1].read_bytes()

    for records in (_records(csv_paths[0]), _records(tsv_paths[0], delimiter="\t")):
        units_row = next(row for row in records if row["metric"] == "pocket_report_units")
        assert json.loads(units_row["value"]) == dict(pockets.units)
        for name, channel in channels.items():
            diagnostic = channel.diagnostics[0]
            row = next(
                item for item in records if item["role"] == f"concordance:{name}" and item["metric"] == diagnostic.code
            )
            assert row["value"] == diagnostic.message
            assert row["status"] == channel.availability.value
            expected_reason = (
                "" if channel.availability is Availability.AVAILABLE else channel.availability.value.replace("_", " ")
            )
            assert row["reason"] == expected_reason
            assert json.loads(row["provenance"]) == channel.to_json()["provenance"]

    pocket_rows = _sheet_rows(xlsx_paths[0], "Pockets")
    units_row = next(row for row in pocket_rows if row["Metric"] == "pocket_report_units")
    assert json.loads(units_row["Value"]) == dict(pockets.units)
    diagnostics = _sheet_rows(xlsx_paths[0], "Diagnostics")
    for name, channel in channels.items():
        diagnostic = channel.diagnostics[0]
        row = next(
            item for item in diagnostics if item["role"] == f"concordance:{name}" and item["metric"] == diagnostic.code
        )
        assert row["status"] == channel.availability.value
        expected_reason = (
            "" if channel.availability is Availability.AVAILABLE else channel.availability.value.replace("_", " ")
        )
        assert (row["reason"] or "") == expected_reason
        assert json.loads(row["provenance"]) == channel.to_json()["provenance"]


def _capture_same_logical_sources(tmp_path: Path) -> tuple[SourceSnapshot, SourceSnapshot, SourceSnapshot]:
    logical = b"ATOM      1  CA  ALA A   1      10.000  10.000  10.000\nEND\n"
    paths = tuple(tmp_path / name for name in ("reference.pdb.gz", "target.pdb.gz", "other.pdb.gz"))
    for path, mtime in zip(paths, (1, 2, 3), strict=True):
        path.write_bytes(gzip.compress(logical, mtime=mtime))
    snapshots = tuple(capture_snapshot(path) for path in paths)
    assert len({item.raw_sha256 for item in snapshots}) == 3
    assert len({item.content_id for item in snapshots}) == 1
    return snapshots


def test_same_content_distinct_raw_snapshots_bind_by_role_and_persist_without_collision(tmp_path: Path) -> None:
    reference, target, conflict = _capture_same_logical_sources(tmp_path)
    report = fixtures._report(reference, target)
    first = serialize_report(report, snapshots=(reference, target))
    assert serialize_report(report, snapshots=(target, reference)) == first
    assert (
        serialize_report(
            report,
            source_paths=(tmp_path / "target.pdb.gz", tmp_path / "reference.pdb.gz"),
        )
        == first
    )
    with pytest.raises(ValueError, match="provenance|role|raw|evidence"):
        serialize_report(report, snapshots=(reference, conflict))
    with pytest.raises(ValueError, match="complete|missing"):
        serialize_report(report, snapshots=(reference,))

    evidence = tmp_path / "evidence"
    serialize_report(report, snapshots=(target, reference), snapshot_dir=evidence)
    reference_path = snapshot_path(reference.content_id, evidence, reference.raw_sha256)
    target_path = snapshot_path(target.content_id, evidence, target.raw_sha256)
    assert reference_path != target_path
    assert hashlib.sha256(reference_path.read_bytes()).hexdigest() == reference.raw_sha256
    assert hashlib.sha256(target_path.read_bytes()).hexdigest() == target.raw_sha256
    assert load_snapshot(reference.content_id, evidence, reference.raw_sha256) == reference
    assert load_snapshot(target.content_id, evidence, target.raw_sha256) == target
    with pytest.raises(ValueError, match="ambiguous|raw"):
        load_snapshot(reference.content_id, evidence)
    assert serialize_report(report, snapshot_dir=evidence) == first

    target_path.write_bytes(b"tampered")
    with pytest.raises(ValueError, match="hash|tamper|snapshot"):
        load_snapshot(target.content_id, evidence, target.raw_sha256)


def test_snapshot_store_preserves_exact_raw_identity_and_unambiguous_legacy_loading(tmp_path: Path) -> None:
    reference, target, _conflict = _capture_same_logical_sources(tmp_path)
    evidence = tmp_path / "evidence"
    store_snapshot(reference, evidence)
    exact = store_snapshot(target, evidence)
    assert store_snapshot(target, evidence) == exact
    renamed = replace(target, display_name="renamed.pdb.gz")
    with pytest.raises(ValueError, match="conflicts"):
        store_snapshot(renamed, evidence)

    legacy_root = tmp_path / "legacy"
    legacy_root.mkdir()
    legacy_path = snapshot_path(reference.content_id, legacy_root)
    legacy_path.write_bytes(reference.decompressed_bytes)
    legacy_path.with_suffix(".json").write_text(
        json.dumps(
            {
                "content_id": reference.content_id,
                "raw_sha256": reference.raw_sha256,
                "decompressed_sha256": reference.decompressed_sha256,
                "display_name": reference.display_name,
                "logical_format": reference.logical_format,
                "is_gzip": reference.is_gzip,
            }
        ),
        encoding="utf-8",
    )
    restored = load_snapshot(reference.content_id, legacy_root)
    assert restored.content_id == reference.content_id
    assert restored.raw_sha256 == reference.raw_sha256

    external_raw = hashlib.sha256(b"external container").hexdigest()
    external = SourceSnapshot(
        b"logical",
        external_raw,
        hashlib.sha256(b"logical").hexdigest(),
        hashlib.sha256(b"logical").hexdigest(),
        "external.pdb",
        "pdb",
        False,
    )
    external_root = tmp_path / "external"
    store_snapshot(external, external_root)
    assert load_snapshot(external.content_id, external_root, external.raw_sha256) == external


def test_snapshot_loading_rejects_invalid_raw_gzip_and_metadata_shapes(tmp_path: Path) -> None:
    logical = b"ATOM\n"
    invalid_raw = b"not gzip"
    snapshot = SourceSnapshot(
        logical,
        hashlib.sha256(invalid_raw).hexdigest(),
        hashlib.sha256(logical).hexdigest(),
        hashlib.sha256(logical).hexdigest(),
        "broken.pdb.gz",
        "pdb",
        True,
        invalid_raw,
    )
    evidence = tmp_path / "evidence"
    store_snapshot(snapshot, evidence)
    with pytest.raises(ValueError, match="gzip|snapshot"):
        load_snapshot(snapshot.content_id, evidence, snapshot.raw_sha256)

    reference = fixtures._snapshot("reference.pdb", b"reference")
    malformed = tmp_path / "malformed"
    store_snapshot(reference, malformed)
    metadata_path = snapshot_path(reference.content_id, malformed).with_suffix(".json")
    original = json.loads(metadata_path.read_text(encoding="utf-8"))
    metadata_path.write_text("{", encoding="utf-8")
    with pytest.raises(ValueError, match="metadata"):
        load_snapshot(reference.content_id, malformed)
    original["storage_encoding"] = "unknown"
    metadata_path.write_text(json.dumps(original), encoding="utf-8")
    with pytest.raises(ValueError, match="metadata"):
        load_snapshot(reference.content_id, malformed)
    with pytest.raises(ValueError, match="missing"):
        load_snapshot("0" * 64, malformed, "1" * 64)


def test_role_binding_rejects_duplicate_snapshots_and_missing_source_paths(tmp_path: Path) -> None:
    reference = fixtures._snapshot("reference.pdb", b"reference")
    target = fixtures._snapshot("target.pdb", b"target")
    report = fixtures._report(reference, target)
    duplicate = replace(reference, display_name="renamed-reference.pdb")
    with pytest.raises(ValueError, match="conflicting"):
        serialize_report(report, snapshots=(reference, duplicate))
    with pytest.raises(FileNotFoundError):
        serialize_report(report, source_paths=(tmp_path / "missing.pdb",))


def test_bounded_read_rejects_same_inode_rewrite_on_first_read(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source = tmp_path / "source.bin"
    source.write_bytes(b"original")
    inode = source.stat().st_ino
    original_read = safe_io.os.read
    first = True

    def rewrite_same_inode(descriptor: int, size: int) -> bytes:
        nonlocal first
        chunk = original_read(descriptor, size)
        if first:
            first = False
            with source.open("r+b", buffering=0) as handle:
                handle.write(b"mutated!")
                os.fsync(handle.fileno())
            assert source.stat().st_ino == inode
        return chunk

    monkeypatch.setattr(safe_io.os, "read", rewrite_same_inode)
    with pytest.raises(ValueError, match="changed while reading|mutated"):
        safe_io.read_bounded_bytes(source, max_bytes=100)


def test_atomic_write_keeps_stable_parent_during_path_swap(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    parent = tmp_path / "stable"
    parent.mkdir()
    target = parent / "artifact.bin"
    moved = tmp_path / "moved"
    original_replace = safe_io.os.replace
    attack_attempted = False
    rename_blocked = False

    def attack_then_replace(source: object, destination: object, *args: object, **kwargs: object) -> None:
        nonlocal attack_attempted, rename_blocked
        if not attack_attempted:
            attack_attempted = True
            try:
                original_replace(parent, moved)
            except OSError:
                rename_blocked = True
            else:
                parent.mkdir()
        original_replace(source, destination, *args, **kwargs)

    monkeypatch.setattr(safe_io.os, "replace", attack_then_replace)
    safe_io.atomic_write_bytes(target, b"payload", max_bytes=100)
    assert attack_attempted
    if rename_blocked:
        assert target.read_bytes() == b"payload"
    else:
        assert (moved / target.name).read_bytes() == b"payload"
        assert not target.exists()


def test_windows_directory_lock_fails_closed_for_missing_parent(tmp_path: Path) -> None:
    if os.name != "nt":
        pytest.skip("Windows handle semantics")
    with pytest.raises(ValueError, match="held safely"):
        safe_io._lock_windows_directories(tmp_path / "missing", label="artifact")


@pytest.mark.parametrize("use_fallback", (False, True))
def test_mutation_residues_and_chain_locators_are_closed_typed_objects(
    monkeypatch: pytest.MonkeyPatch, use_fallback: bool
) -> None:
    report, _reference, _target, _event = round2._mutation_report()
    payload = report_to_dict(report)
    payload["reference_selection"]["chain_locators"] = [
        {"author_chain_id": "A", "label_chain_id": "X", "entity_id": "1"}
    ]
    validate_analysis_report_payload(payload)

    if use_fallback:
        real_import = schema_module.importlib.import_module

        def missing_jsonschema(name: str) -> object:
            if name == "jsonschema":
                raise ImportError("optional dependency absent")
            return real_import(name)

        monkeypatch.setattr(schema_module.importlib, "import_module", missing_jsonschema)

    bad_mutation = json.loads(json.dumps(payload))
    bad_mutation["pockets"]["comparisons"][0]["associated_mutations"][0]["reference"]["unknown"] = True
    with pytest.raises(ValueError, match="unknown|schema"):
        validate_analysis_report_payload(bad_mutation)

    bad_locator = json.loads(json.dumps(payload))
    bad_locator["reference_selection"]["chain_locators"][0]["unknown"] = True
    with pytest.raises(ValueError, match="unknown|schema"):
        validate_analysis_report_payload(bad_locator)

    empty_locator = json.loads(json.dumps(payload))
    empty_locator["reference_selection"]["chain_locators"] = [
        {"author_chain_id": None, "label_chain_id": None, "entity_id": None}
    ]
    with pytest.raises(ValueError, match="locator|chain|schema"):
        validate_analysis_report_payload(empty_locator)

    invalid_entity = json.loads(json.dumps(payload))
    invalid_entity["reference_selection"]["chain_locators"] = [
        {"author_chain_id": "A", "label_chain_id": None, "entity_id": ""}
    ]
    with pytest.raises(ValueError, match="entity|locator|schema"):
        validate_analysis_report_payload(invalid_entity)
