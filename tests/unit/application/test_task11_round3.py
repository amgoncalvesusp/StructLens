from __future__ import annotations

import importlib
from dataclasses import replace
from pathlib import Path

import pytest
from jsonschema import Draft202012Validator
from openpyxl import load_workbook

import structlens.application.report_exports as report_exports
import structlens.core.reports.safe_io as safe_io
import structlens.core.reports.schema as schema_module
from structlens.application.export_service import export_report_csv, export_report_tsv, export_report_xlsx
from structlens.application.report_serialization import (
    deserialize_report,
    report_to_dict,
    serialize_report,
)
from structlens.core.evidence import Availability, Diagnostic, DiagnosticSeverity
from structlens.core.parsing import SourceSnapshot, capture_snapshot
from structlens.core.provenance import MethodProvenance
from structlens.core.reports.schema import load_analysis_report_schema, validate_analysis_report_payload

fixtures = importlib.import_module("tests.unit.application.test_task11_round1")


def test_bundled_schema_has_valid_and_resolvable_local_references() -> None:
    """The packaged schema must not ship a dangling local ``$ref``."""

    bundled = load_analysis_report_schema()
    Draft202012Validator.check_schema(bundled)
    definitions = bundled.get("$defs", {})
    missing: list[str] = []

    def visit(node: object) -> None:
        if isinstance(node, dict):
            reference = node.get("$ref")
            if isinstance(reference, str) and reference.startswith("#/$defs/"):
                target = reference.removeprefix("#/$defs/")
                if target not in definitions:
                    missing.append(reference)
            for child in node.values():
                visit(child)
        elif isinstance(node, list):
            for child in node:
                visit(child)

    visit(bundled)
    assert missing == []


def _rich_pocket_report() -> tuple[object, SourceSnapshot, SourceSnapshot]:
    reference = fixtures._snapshot("reference.pdb", b"reference")
    target = fixtures._snapshot("target.pdb", b"target")
    report = fixtures._report(reference, target, pockets=fixtures._pocket_report(reference, target))
    assert report.pockets is not None and report.pockets.matching is not None
    matching = replace(
        report.pockets.matching,
        diagnostics=(Diagnostic("MATCH_GLOBAL", DiagnosticSeverity.WARNING, "ambiguous global assignment"),),
    )
    provenance = MethodProvenance(
        "pocket.pipeline",
        "2",
        parameters={"detector": "alpha-sphere"},
        units={"volume": "angstrom^3"},
        input_hashes={"source": reference.content_id},
    )
    return replace(report, pockets=replace(report.pockets, matching=matching, provenance=provenance)), reference, target


def test_all_delimited_and_xlsx_exports_retain_pocket_identities(tmp_path: Path) -> None:
    report, reference, target = _rich_pocket_report()
    rows = report_exports._rows(report)
    candidate_id = report.pockets.volumes[0].result.candidate_id
    assert any(row["metric"] == "volume_candidate_id" and row["value"] == candidate_id for row in rows)
    assert any(
        row["section"] == "diagnostics" and row["role"] == "matching" and row["metric"] == "MATCH_GLOBAL"
        for row in rows
    )
    assert any(
        row["section"] == "provenance"
        and row["metric"] == "pocket_artifact_id"
        and row["value"] == report.pockets.provenance.artifact_id
        for row in rows
    )

    csv_path = tmp_path / "report.csv"
    tsv_path = tmp_path / "report.tsv"
    xlsx_path = tmp_path / "report.xlsx"
    export_report_csv(report, csv_path, snapshots=(reference, target))
    export_report_tsv(report, tsv_path, snapshots=(reference, target))
    export_report_xlsx(report, xlsx_path, snapshots=(reference, target))
    for path in (csv_path, tsv_path):
        text = path.read_text(encoding="utf-8")
        assert candidate_id in text
        assert "MATCH_GLOBAL" in text
        assert report.pockets.provenance.artifact_id in text
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    cells = [cell.value for row in workbook["Pockets"].iter_rows() for cell in row]
    cells.extend(cell.value for row in workbook["Matches"].iter_rows() for cell in row)
    cells.extend(cell.value for row in workbook["Provenance"].iter_rows() for cell in row)
    assert candidate_id in cells
    assert "MATCH_GLOBAL" in cells
    assert report.pockets.provenance.artifact_id in cells


def test_capture_snapshot_reads_source_once_through_safe_handle(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "source.pdb"
    source.write_bytes(b"ATOM source")

    def forbidden_path_open(*_args: object, **_kwargs: object) -> object:
        raise AssertionError("capture must not reopen the source by pathname")

    monkeypatch.setattr(Path, "open", forbidden_path_open)
    snapshot = capture_snapshot(source)
    assert snapshot.decompressed_bytes == b"ATOM source"


def test_safe_io_fails_closed_when_ancestor_changes_before_read(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    source = tmp_path / "source.bin"
    source.write_bytes(b"payload")
    original_check = safe_io._assert_ancestor_state
    calls = 0

    def race_check(*args: object, **kwargs: object) -> None:
        nonlocal calls
        calls += 1
        if calls >= 2:
            raise ValueError("file ancestor changed while reading")
        original_check(*args, **kwargs)

    monkeypatch.setattr(safe_io, "_assert_ancestor_state", race_check)
    with pytest.raises(ValueError, match="ancestor changed"):
        safe_io.read_bounded_bytes(source, max_bytes=100)


def test_atomic_write_fails_closed_when_ancestor_changes_before_replace(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    target = tmp_path / "artifact.bin"
    original_check = safe_io._assert_ancestor_state
    calls = 0

    def race_check(*args: object, **kwargs: object) -> None:
        nonlocal calls
        calls += 1
        if calls >= 2:
            raise ValueError("file ancestor changed before replacement")
        original_check(*args, **kwargs)

    monkeypatch.setattr(safe_io, "_assert_ancestor_state", race_check)
    with pytest.raises(ValueError, match="ancestor changed"):
        safe_io.atomic_write_bytes(target, b"payload", max_bytes=100)
    assert not target.exists()


@pytest.mark.parametrize("use_fallback", (False, True))
def test_nested_report_objects_are_closed_world(monkeypatch: pytest.MonkeyPatch, use_fallback: bool) -> None:
    reference = fixtures._snapshot("reference.pdb", b"reference")
    target = fixtures._snapshot("target.pdb", b"target")
    payload = report_to_dict(fixtures._report(reference, target, pockets=fixtures._pocket_report(reference, target)))
    payload["analysis"] = {"unknown": 1}
    if use_fallback:
        real_import = schema_module.importlib.import_module

        def missing_jsonschema(name: str) -> object:
            if name == "jsonschema":
                raise ImportError("optional dependency absent")
            return real_import(name)

        monkeypatch.setattr(schema_module.importlib, "import_module", missing_jsonschema)
    with pytest.raises(ValueError, match="unknown|schema"):
        validate_analysis_report_payload(payload)

    payload = report_to_dict(fixtures._report(reference, target, pockets=fixtures._pocket_report(reference, target)))
    payload["evidence_cards"] = [{"unknown": 1}]
    with pytest.raises(ValueError, match="unknown|schema"):
        validate_analysis_report_payload(payload)

    payload = report_to_dict(fixtures._report(reference, target, pockets=fixtures._pocket_report(reference, target)))
    payload["pockets"]["detections"][0]["unknown"] = 1
    with pytest.raises(ValueError, match="unknown|schema"):
        validate_analysis_report_payload(payload)


def test_pocket_aggregate_availability_rejects_child_contradiction() -> None:
    reference = fixtures._snapshot("reference.pdb", b"reference")
    pockets = fixtures._pocket_report(reference)
    with pytest.raises(ValueError, match="availability|coherent"):
        replace(pockets, availability=Availability.INVALID_INPUT)


def test_identical_content_snapshots_bind_to_both_roles() -> None:
    snapshot = fixtures._snapshot("structure.pdb", b"same content")
    report = fixtures._report(snapshot, snapshot)
    restored = deserialize_report(
        serialize_report(report, snapshots=(snapshot, snapshot)), snapshots=(snapshot, snapshot)
    )
    assert restored.reference_selection.content_id == restored.target_selection.content_id == snapshot.content_id


def test_surface_method_provenance_round_trips_as_typed_value() -> None:
    report, reference, target = _rich_pocket_report()
    assert report.pockets is not None
    original = MethodProvenance("surface.method", "1", input_hashes={"source": reference.content_id})
    comparison = replace(report.pockets.comparisons[0], reference_surface_provenance=original)
    rich = replace(report, pockets=replace(report.pockets, comparisons=(comparison, *report.pockets.comparisons[1:])))
    restored = deserialize_report(serialize_report(rich, snapshots=(reference, target)), snapshots=(reference, target))
    assert restored.pockets is not None
    assert restored.pockets.comparisons[0].reference_surface_provenance == original


@pytest.mark.parametrize("status", (Availability.INVALID_INPUT, Availability.NUMERICAL_FAILURE))
def test_top_level_diagnostic_uses_native_report_status(status: Availability) -> None:
    reference = fixtures._snapshot("reference.pdb", b"reference")
    target = fixtures._snapshot("target.pdb", b"target")
    report = fixtures._report(reference, target)
    rich = replace(
        report,
        availability=replace(report.availability, analysis=status),
        diagnostics=(Diagnostic("REPORT_STATE", DiagnosticSeverity.ERROR, "report state"),),
    )
    row = next(item for item in report_exports._rows(rich) if item["metric"] == "REPORT_STATE")
    assert row["status"] == status.value
    assert row["reason"] == status.value.replace("_", " ")
